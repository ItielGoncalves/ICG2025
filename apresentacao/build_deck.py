#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Gera a apresentacao institucional "Avanco da Automacao na Inpasa".
Deck neutro de fornecedor, com dados reais (estado das malhas + piloto Dourados).
Saida: Visao_Inpasa_Automacao.pptx
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.chart.data import CategoryChartData
from pptx.oxml.ns import qn

# ----- Paleta -----
NAVY   = RGBColor(0x12, 0x2A, 0x4A)   # azul profundo Inpasa
BLUE   = RGBColor(0x1F, 0x5C, 0xA8)
BLUE2  = RGBColor(0x3E, 0x7C, 0xC4)
AMBER  = RGBColor(0xF2, 0xA0, 0x33)   # chama Inpasa
GREEN  = RGBColor(0x2F, 0x9E, 0x44)
RED    = RGBColor(0xC8, 0x47, 0x3B)
YELLOW = RGBColor(0xE2, 0xB0, 0x33)
LIGHT  = RGBColor(0xF4, 0xF6, 0xF9)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
INK    = RGBColor(0x1B, 0x2A, 0x41)
GRAY   = RGBColor(0x5B, 0x6B, 0x7F)
LINE   = RGBColor(0xD7, 0xDE, 0xE7)

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]

def slide():
    return prs.slides.add_slide(BLANK)

def bg(s, color):
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = color

def rect(s, x, y, w, h, fill=None, line=None, line_w=1.0, shape=MSO_SHAPE.RECTANGLE, shadow=False):
    sp = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line; sp.line.width = Pt(line_w)
    sp.shadow.inherit = False
    if shadow:
        spPr = sp._element.spPr
        ef = spPr.makeelement(qn('a:effectLst'), {})
        sh = ef.makeelement(qn('a:outerShdw'),
             {'blurRad':'80000','dist':'30000','dir':'5400000','rotWithShape':'0'})
        clr = sh.makeelement(qn('a:srgbClr'), {'val':'1B2A41'})
        a = clr.makeelement(qn('a:alpha'), {'val':'20000'})
        clr.append(a); sh.append(clr); ef.append(sh)
        spPr.append(ef)   # effectLst must be the last child of spPr
    return sp

def txt(s, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
        space_after=4, line_spacing=1.0):
    """runs: list of paragraphs; each paragraph = list of (text,size,color,bold,italic)."""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Pt(0); tf.margin_top = tf.margin_bottom = Pt(0)
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(space_after); p.space_before = Pt(0)
        p.line_spacing = line_spacing
        for (t, sz, col, bold, ital) in para:
            r = p.add_run(); r.text = t
            r.font.size = Pt(sz); r.font.color.rgb = col
            r.font.bold = bold; r.font.italic = ital
            r.font.name = "Calibri"
    return tb

def notes(s, text):
    s.notes_slide.notes_text_frame.text = text

def footer(s, page, dark=False):
    c = WHITE if dark else GRAY
    txt(s, 0.55, 7.02, 8, 0.3,
        [[("INPASA  ·  Avanço da Automação — Visão e Jornada", 9, c, False, False)]])
    txt(s, 11.6, 7.02, 1.2, 0.3,
        [[(f"{page:02d}", 9, AMBER, True, False)]], align=PP_ALIGN.RIGHT)

def kicker(s, label):
    rect(s, 0.55, 0.62, 0.16, 0.46, fill=AMBER)
    txt(s, 0.82, 0.6, 11, 0.5, [[(label, 13, BLUE, True, False)]])

def title(s, text, y=1.05, size=30, color=NAVY, w=12.2):
    txt(s, 0.82, y, w, 1.0, [[(text, size, color, True, False)]])

# ===================================================================
# SLIDE 1 — Capa
# ===================================================================
s = slide(); bg(s, NAVY)
rect(s, 0, 0, 13.333, 0.28, fill=AMBER)
rect(s, 0, 7.22, 13.333, 0.28, fill=BLUE)
# selo
rect(s, 0.9, 1.5, 0.9, 0.9, fill=AMBER, shape=MSO_SHAPE.OVAL)
txt(s, 0.9, 1.5, 0.9, 0.9, [[("IA", 26, NAVY, True, False)]],
    align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
txt(s, 0.9, 3.05, 11.6, 2.2, [
    [("Avanço da Automação", 52, WHITE, True, False)],
    [("na Inpasa", 52, AMBER, True, False)],
], line_spacing=1.0)
rect(s, 0.95, 5.35, 3.2, 0.05, fill=BLUE2)
txt(s, 0.9, 5.55, 11.5, 0.6, [[("De controlar a planta a governar a planta", 20, LIGHT, False, True)]])
txt(s, 0.9, 6.45, 11.5, 0.5,
    [[("Apresentação institucional  ·  Liderança e Parceiros de Automação Industrial", 13, RGBColor(0xAEB9C8 // 0x10000, (0xAEB9C8 // 0x100) % 0x100, 0xAEB9C8 % 0x100), False, False)]])
notes(s, "Bom dia a todos. Em nome da Inpasa, agradeço a recepção. Nos próximos minutos quero "
          "compartilhar não uma lista de tecnologias, mas a visão que está guiando o avanço da "
          "automação na Inpasa — onde estamos hoje, de forma medida, e para onde estamos indo. "
          "É essa visão que queremos construir junto com nossos parceiros.")

# ===================================================================
# SLIDE 2 — A virada de chave (tese)
# ===================================================================
s = slide(); bg(s, LIGHT)
kicker(s, "A VIRADA DE CHAVE")
title(s, "Automação deixou de ser controlar a planta.")
txt(s, 0.82, 1.62, 12.0, 1.0,
    [[("Passou a ", 30, NAVY, True, False), ("enxergar e governar", 30, AMBER, True, False),
      (" o estado dela em tempo real.", 30, NAVY, True, False)]])
txt(s, 0.82, 2.95, 11.6, 0.6,
    [[("O operador deixa de atuar o tempo todo e passa a ser o ", 16, GRAY, False, False),
      ("vigia", 16, BLUE, True, False),
      (". O sistema opera — e audita a si mesmo.", 16, GRAY, False, False)]])
pilares = [("Controle", "Malhas estáveis, automáticas e auditadas", GREEN),
           ("Gestão de Alarmes", "Foco no que exige ação do operador", BLUE),
           ("Manutenção Preditiva", "Agir antes da falha, preservando ativos", AMBER)]
pw, gap = 3.78, 0.34
x0 = 0.82
for i,(t,d,c) in enumerate(pilares):
    x = x0 + i*(pw+gap)
    rect(s, x, 4.1, pw, 2.05, fill=WHITE, line=LINE, line_w=1.0, shadow=True)
    rect(s, x, 4.1, pw, 0.16, fill=c)
    txt(s, x+0.32, 4.5, pw-0.6, 0.6, [[(f"0{i+1}", 16, c, True, False)]])
    txt(s, x+0.32, 4.95, pw-0.6, 0.6, [[(t, 19, NAVY, True, False)]])
    txt(s, x+0.32, 5.5, pw-0.6, 0.7, [[(d, 13, GRAY, False, False)]])
footer(s, 2)
notes(s, "A nossa tese é simples. Por muito tempo, automação significou controlar a planta — "
          "manter as variáveis no setpoint. Hoje, o salto que perseguimos é outro: enxergar e "
          "governar o estado da planta em tempo real. O operador deixa de ser quem atua o tempo "
          "todo e passa a ser o vigia; o sistema opera e, tão importante quanto, audita a si "
          "mesmo. Tudo se apoia em três pilares: controle, gestão de alarmes e manutenção preditiva.")

# ===================================================================
# SLIDE 3 — Onde estamos hoje (evidencia medida) — DASHBOARD
# ===================================================================
s = slide(); bg(s, LIGHT)
kicker(s, "ONDE ESTAMOS HOJE  ·  EVIDÊNCIA MEDIDA")
title(s, "Estado das malhas de controle — últimos 15 dias", size=26)

def kpi_row(s, y, header, tiles):
    txt(s, 0.55, y, 12.3, 0.28, [[(header, 12, NAVY, True, False)]])
    n = len(tiles); g = 0.12; x0 = 0.55
    tw = (12.33 - (n-1)*g)/n
    for i,(lab,val,col) in enumerate(tiles):
        x = x0 + i*(tw+g)
        rect(s, x, y+0.32, tw, 0.92, fill=col)
        txt(s, x, y+0.40, tw, 0.46, [[(str(val), 26, WHITE, True, False)]],
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        txt(s, x, y+0.96, tw, 0.26, [[(lab, 9.5, WHITE, True, False)]],
            align=PP_ALIGN.CENTER)

util = [("AUTO",87,GREEN),("MANUAL",6,RED),("PROGRAMA",82,GREEN),
        ("OPERADOR",54,YELLOW),("BYPASS",1,RED),("MANUTENÇÃO",8,YELLOW)]
proc = [("AUTO",147,GREEN),("MANUAL",22,RED),("PROGRAMA",125,GREEN),
        ("OPERADOR",27,YELLOW),("BYPASS",0,GREEN),("MANUTENÇÃO",1,YELLOW)]
kpi_row(s, 1.55, "MALHAS DE UTILIDADES", util)
kpi_row(s, 2.92, "MALHAS DE PROCESSO", proc)

# Donuts
def donut(s, x, y, cats, vals, colors, caption):
    cd = CategoryChartData(); cd.categories = cats
    cd.add_series("s", vals)
    gf = s.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, Inches(x), Inches(y),
                            Inches(2.5), Inches(1.95), cd)
    ch = gf.chart
    ch.has_legend = False; ch.has_title = False
    plot = ch.plots[0]; plot.donut_hole_size = 62
    ser = plot.series[0]
    for idx, col in enumerate(colors):
        pt = ser.points[idx]
        pt.format.fill.solid(); pt.format.fill.fore_color.rgb = col
        pt.format.line.color.rgb = WHITE; pt.format.line.width = Pt(1.5)
    dl = plot.data_labels
    dl.show_percentage = True; dl.show_value = False
    dl.number_format = '0%'; dl.number_format_is_linked = False
    dl.font.size = Pt(10); dl.font.bold = True; dl.font.color.rgb = WHITE
    dl.position = XL_LABEL_POSITION.CENTER
    txt(s, x-0.25, y+1.95, 3.0, 0.5, [[(caption, 10.5, NAVY, True, False)]],
        align=PP_ALIGN.CENTER)

dy = 4.45
donut(s, 0.65, dy, ["AUTO","MANUAL"], [94,6], [GREEN,RED], "Utilidades — Auto × Manual")
donut(s, 3.75, dy, ["PROGRAMA","OPERADOR"], [60,40], [GREEN,YELLOW], "Utilidades — Programa × Operador")
donut(s, 6.95, dy, ["AUTO","MANUAL"], [87,13], [GREEN,RED], "Processo — Auto × Manual")
donut(s, 10.05, dy, ["PROGRAMA","OPERADOR"], [82,18], [GREEN,YELLOW], "Processo — Programa × Operador")

footer(s, 3)
notes(s, "E não falo de aspiração — falo de dado. Este é o estado real das nossas malhas de "
          "controle nos últimos 15 dias. 94% das malhas de utilidades e 87% das de processo "
          "operando em automático. Mas o mais importante não é só o número alto: é que hoje "
          "enxergamos cada exceção — cada malha em manual, cada interlock em bypass, cada "
          "variável em modo operador. O que antes era invisível agora é medido e auditado. "
          "Isso é governança de automação. (Se perguntarem do método: o estado é o visto pela "
          "HMI; mudanças por lógica de PLC ainda não são auditadas — transparência reforça "
          "credibilidade.)")

# ===================================================================
# SLIDE 4 — Gestao de ativos: piloto Dourados
# ===================================================================
s = slide(); bg(s, LIGHT)
kicker(s, "GESTÃO DE ATIVOS  ·  PILOTO EM DOURADOS")
title(s, "Da malha ao ativo: visibilidade e disciplina operacional")
# numeros grandes
for i,(num,lab) in enumerate([("5","objetos simulados\n(unidade piloto)"),
                              ("48","objetos simulados\n(visão consolidada)")]):
    x = 0.82 + i*3.0
    rect(s, x, 2.0, 2.7, 1.9, fill=NAVY, shadow=True)
    txt(s, x, 2.15, 2.7, 1.0, [[(num, 54, AMBER, True, False)]],
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    txt(s, x, 3.25, 2.7, 0.6, [[(lab.replace("\n"," "), 12, WHITE, False, False)]],
        align=PP_ALIGN.CENTER)
caps = [("Equipamentos em simulação", "monitorados em tempo real"),
        ("Ativos inoperantes e fora de operação", "identificados automaticamente"),
        ("Auditoria de usuário", "quem fez o quê, em qual estação"),
        ("Tempo em simulação", "métrica inédita de disciplina operacional")]
yx = 2.0
for i,(t,d) in enumerate(caps):
    y = yx + i*0.52
    rect(s, 7.0, y+0.06, 0.16, 0.16, fill=AMBER, shape=MSO_SHAPE.OVAL)
    txt(s, 7.3, y, 5.6, 0.5,
        [[(t+"  ", 14, NAVY, True, False), (d, 12.5, GRAY, False, False)]])
rect(s, 0.82, 4.35, 11.7, 0.9, fill=RGBColor(0xE9,0xF0,0xF8), line=BLUE2, line_w=1.0)
txt(s, 1.1, 4.45, 11.2, 0.7,
    [[("Plano:  ", 14, BLUE, True, False),
      ("validar em Dourados e ", 14, INK, False, False),
      ("padronizar para todo o Grupo Inpasa", 14, NAVY, True, False),
      (" — gestão unificada e visibilidade operacional dos ativos.", 14, INK, False, False)]],
    anchor=MSO_ANCHOR.MIDDLE)
footer(s, 4)
notes(s, "Demos um passo além do controle: a gestão de ativos. Iniciamos em Dourados um piloto "
          "que monitora equipamentos em simulação, identifica ativos inoperantes e equipamentos "
          "fora de operação, e audita quem fez o quê e por quanto tempo. Pela primeira vez temos "
          "visibilidade de quanto tempo um ativo permanece simulado — uma métrica de disciplina "
          "operacional que não tínhamos. O plano é validar em Dourados e padronizar para todo o "
          "Grupo Inpasa.")

# ===================================================================
# SLIDE 5 — Operacao Observatoria
# ===================================================================
s = slide(); bg(s, NAVY)
rect(s, 0.55, 0.62, 0.16, 0.46, fill=AMBER)
txt(s, 0.82, 0.6, 11, 0.5, [[("CONCEITO  ·  OPERAÇÃO OBSERVATÓRIA", 13, AMBER, True, False)]])
txt(s, 0.82, 1.15, 12, 1.0, [[("O operador supervisiona. O sistema opera.", 32, WHITE, True, False)]])
# dois cards conceito
cards = [("Operador → Vigia",
          "Supervisiona, decide sobre exceções e direciona. Sai do controle manual repetitivo."),
         ("Sistema → Opera e audita",
          "Mantém as malhas em automático e registra cada desvio do estado normal.")]
for i,(t,d) in enumerate(cards):
    x = 0.82 + i*6.05
    rect(s, x, 2.5, 5.7, 1.5, fill=RGBColor(0x1B,0x3A,0x5E), line=BLUE2, line_w=1.0)
    txt(s, x+0.35, 2.7, 5.0, 0.5, [[(t, 18, AMBER, True, False)]])
    txt(s, x+0.35, 3.25, 5.0, 0.7, [[(d, 13.5, LIGHT, False, False)]])
rect(s, 0.82, 4.45, 11.7, 1.55, fill=RGBColor(0x14,0x31,0x52), line=AMBER, line_w=1.25)
txt(s, 1.15, 4.65, 11.0, 0.5,
    [[("As malhas de controle viram sensores de saúde.", 19, WHITE, True, False)]])
txt(s, 1.15, 5.25, 11.0, 0.7,
    [[("Desvios no comportamento de uma malha antecipam a falha antes que ela vire um evento — "
       "detecção preditiva nascendo de dentro do próprio controle.", 14, LIGHT, False, False)]])
footer(s, 5, dark=True)
notes(s, "Esse caminho nos leva ao conceito que chamamos de operação observatória. O operador "
          "supervisiona; o sistema opera. E vamos além: as próprias malhas de controle viram "
          "sensores de saúde — desvios no comportamento de uma malha antecipam uma falha antes "
          "que ela vire um evento. É a manutenção preditiva nascendo de dentro do próprio controle.")

# ===================================================================
# SLIDE 6 — Arquitetura de dados (DataOps)
# ===================================================================
s = slide(); bg(s, LIGHT)
kicker(s, "A BASE  ·  ARQUITETURA DE DADOS")
title(s, "Dados confiáveis e organizados são a fundação de tudo")
camadas = [("Coleta","De controladores e instrumentos",BLUE),
           ("Historiação","Série temporal confiável",BLUE2),
           ("Contextualização","Dado com significado operacional",AMBER),
           ("Análise","Insight, comparação e decisão",GREEN)]
bw, gp = 2.7, 0.45
x0 = 0.82
for i,(t,d,c) in enumerate(camadas):
    x = x0 + i*(bw+gp)
    rect(s, x, 2.2, bw, 1.5, fill=WHITE, line=LINE, shadow=True)
    rect(s, x, 2.2, 0.14, 1.5, fill=c)
    txt(s, x+0.32, 2.42, bw-0.4, 0.5, [[(t, 17, NAVY, True, False)]])
    txt(s, x+0.32, 2.95, bw-0.4, 0.6, [[(d, 12, GRAY, False, False)]])
    if i < 3:
        txt(s, x+bw, 2.55, 0.45, 0.8, [[("➜", 22, AMBER, True, False)]], align=PP_ALIGN.CENTER)
# dois destaques
dest = [("Varredura automática dos controladores",
         "Padroniza a informação em escala, sem trabalho manual."),
        ("Benchmarking entre plantas",
         "Saber qual unidade opera melhor — e por quê — e levar a melhor prática às demais.")]
for i,(t,d) in enumerate(dest):
    x = 0.82 + i*6.05
    rect(s, x, 4.35, 5.7, 1.55, fill=RGBColor(0xE9,0xF0,0xF8), line=BLUE2, line_w=1.0)
    txt(s, x+0.35, 4.55, 5.0, 0.5, [[(t, 16, BLUE, True, False)]])
    txt(s, x+0.35, 5.1, 5.0, 0.7, [[(d, 13, INK, False, False)]])
footer(s, 6)
notes(s, "Nada disso se sustenta sem dados confiáveis e organizados. Estamos estruturando a "
          "arquitetura de dados operacionais da Inpasa — coleta, historiação e análise — com "
          "varredura automática dos controladores para padronizar a informação. Isso nos abre "
          "algo poderoso: comparar plantas entre si, o benchmarking interno. Saber qual unidade "
          "opera melhor, e por quê, e levar essa melhor prática para as demais.")

# ===================================================================
# SLIDE 7 — IA no controle + manutencao preditiva
# ===================================================================
s = slide(); bg(s, LIGHT)
kicker(s, "INTELIGÊNCIA  ·  IA NO CONTROLE E MANUTENÇÃO")
title(s, "IA onde ela gera valor: controlar melhor e falhar menos")
frentes = [("Controle avançado e adaptativo",
            "Ajusta-se ao processo e reduz variabilidade.", GREEN),
           ("IA que aprende e recomenda",
            "Predição de resultados e apoio à decisão do operador.", BLUE),
           ("Manutenção preditiva de verdade",
            "Do reativo (corretiva/preventiva) ao agir antes da falha.", AMBER)]
pw, gap = 3.78, 0.34
for i,(t,d,c) in enumerate(frentes):
    x = 0.82 + i*(pw+gap)
    rect(s, x, 2.2, pw, 2.0, fill=WHITE, line=LINE, shadow=True)
    rect(s, x, 2.2, pw, 0.16, fill=c)
    txt(s, x+0.32, 2.55, pw-0.6, 0.7, [[(t, 16.5, NAVY, True, False)]])
    txt(s, x+0.32, 3.4, pw-0.6, 0.7, [[(d, 12.5, GRAY, False, False)]])
rect(s, 0.82, 4.7, 11.7, 1.2, fill=NAVY)
res = "Menos paradas    ·    Mais confiabilidade    ·    Ativos preservados    ·    Operação mais autônoma"
txt(s, 0.82, 4.7, 11.7, 1.2, [[(res, 17, WHITE, True, False)]],
    align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
footer(s, 7)
notes(s, "Com a base de dados pronta, a inteligência artificial entra onde ela realmente gera "
          "valor: controle avançado que se adapta ao processo, predição de resultados e "
          "recomendações para o operador. E na manutenção, saímos do ciclo caro de corretiva e "
          "preventiva para a preditiva de verdade — agir antes da falha, preservando ativo e "
          "produção. Menos paradas, mais confiabilidade.")

# ===================================================================
# SLIDE 8 — COI
# ===================================================================
s = slide(); bg(s, NAVY)
rect(s, 0.55, 0.62, 0.16, 0.46, fill=AMBER)
txt(s, 0.82, 0.6, 11, 0.5, [[("VISÃO  ·  CENTRO DE OPERAÇÕES INTEGRADO (COI)", 13, AMBER, True, False)]])
txt(s, 0.82, 1.15, 12, 1.0,
    [[("Uma só sala de inteligência para todo o Grupo", 30, WHITE, True, False)]])
itens = [("Robusto e escalável","Cresce com o Grupo, unidade a unidade."),
         ("Centraliza decisão, não só telas","Visão única e autonomia operacional."),
         ("Otimização de mão de obra","Especialistas focados onde agregam mais.")]
for i,(t,d) in enumerate(itens):
    x = 0.82 + i*3.95
    rect(s, x, 2.5, 3.7, 1.7, fill=RGBColor(0x1B,0x3A,0x5E), line=BLUE2, line_w=1.0)
    txt(s, x+0.3, 2.7, 3.2, 0.7, [[(t, 16, AMBER, True, False)]])
    txt(s, x+0.3, 3.45, 3.2, 0.7, [[(d, 12.5, LIGHT, False, False)]])
txt(s, 0.82, 4.6, 11.7, 0.9,
    [[("Inspirado em modelos já consolidados em indústrias de grande porte — ", 14, LIGHT, False, True),
      ("trazido para a realidade do agronegócio da Inpasa.", 14, WHITE, True, True)]])
footer(s, 8, dark=True)
notes(s, "Tudo converge para o Centro de Operações Integrado. A visão é um COI robusto e "
          "escalável, que integre as unidades do Grupo numa só sala de inteligência — inspirado "
          "em modelos já consolidados em outras indústrias de grande porte, como a mineração. "
          "Mais do que centralizar telas, é centralizar decisão: autonomia operacional, "
          "otimização de mão de obra especializada e uma visão única do Grupo Inpasa.")

# ===================================================================
# SLIDE 9 — Roadmap
# ===================================================================
s = slide(); bg(s, LIGHT)
kicker(s, "ROADMAP  ·  UM CAMINHO JÁ EM MOVIMENTO")
title(s, "Incremental, medido e em execução")
etapas = [("HOJE","Malhas medidas e auditadas + piloto de ativos em Dourados", GREEN),
          ("CURTO","Padronização de dados + gestão de ativos em todo o Grupo", BLUE),
          ("MÉDIO","IA aplicada ao controle e à manutenção preditiva", AMBER),
          ("VISÃO","Centro de Operações Integrado e operação autônoma", NAVY)]
# linha base
rect(s, 1.0, 3.55, 11.3, 0.06, fill=LINE)
bw, gp = 2.7, 0.45; x0 = 0.82
for i,(fase,desc,c) in enumerate(etapas):
    x = x0 + i*(bw+gp)
    cx = x + bw/2
    rect(s, cx-0.16, 3.42, 0.32, 0.32, fill=c, shape=MSO_SHAPE.OVAL)
    rect(s, x, 2.0, bw, 1.15, fill=WHITE, line=LINE, shadow=True)
    rect(s, x, 2.0, bw, 0.5, fill=c)
    txt(s, x, 2.06, bw, 0.42, [[(fase, 15, WHITE, True, False)]],
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    txt(s, x+0.2, 2.58, bw-0.4, 0.55, [[(desc, 11.5, INK, False, False)]], align=PP_ALIGN.CENTER)
    txt(s, x, 3.9, bw, 0.4, [[(["Em andamento","Próximo","A seguir","Destino"][i], 11, c, True, False)]],
        align=PP_ALIGN.CENTER)
rect(s, 0.82, 5.0, 11.7, 1.0, fill=RGBColor(0xE9,0xF0,0xF8), line=BLUE2)
txt(s, 0.82, 5.0, 11.7, 1.0,
    [[("Cada etapa entrega valor por si — e prepara a seguinte.", 16, BLUE, True, False)]],
    align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
footer(s, 9)
notes(s, "O caminho é incremental e já está em movimento: hoje, malhas medidas e o piloto de "
          "Dourados; a seguir, padronização de dados e gestão de ativos em todo o Grupo; depois, "
          "IA aplicada ao controle e à manutenção preditiva; e, como destino, o COI e a operação "
          "autônoma.")

# ===================================================================
# SLIDE 10 — Fechamento
# ===================================================================
s = slide(); bg(s, NAVY)
rect(s, 0, 0, 13.333, 0.28, fill=AMBER)
rect(s, 0, 7.22, 13.333, 0.28, fill=BLUE)
txt(s, 0.9, 2.2, 11.5, 1.6,
    [[("A Inpasa sabe", 46, WHITE, True, False)],
     [("aonde quer chegar.", 46, AMBER, True, False)]], line_spacing=1.0)
rect(s, 0.95, 4.35, 3.2, 0.05, fill=BLUE2)
txt(s, 0.9, 4.6, 11.3, 1.0,
    [[("Damos os passos com método e medição. Buscamos parceiros que ", 16, LIGHT, False, False),
      ("acelerem essa visão", 16, WHITE, True, False),
      (" — com tecnologia robusta, escalável e que respeite a jornada que já construímos.", 16, LIGHT, False, False)]])
txt(s, 0.9, 6.0, 11.3, 0.6, [[("Obrigado.", 22, AMBER, True, False)]])
notes(s, "Encerro com uma mensagem clara: a Inpasa sabe aonde quer chegar, e tem dado os passos "
          "com método e medição. O que buscamos em nossos parceiros é quem nos ajude a acelerar "
          "essa visão — com tecnologia robusta, escalável e que respeite a jornada que já "
          "construímos. Conto com vocês nessa caminhada. Obrigado.")

out = "/home/user/ICG2025/apresentacao/Visao_Inpasa_Automacao.pptx"
prs.save(out)
print("OK ->", out, "| slides:", len(prs.slides._sldIdLst))
