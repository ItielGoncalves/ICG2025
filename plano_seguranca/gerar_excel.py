# -*- coding: utf-8 -*-
"""
Gera o Excel consolidado do Plano de Acao de Seguranca (INPASA).
Abas:
  (1) Matriz de Risco  - transportadores por area e por unidade
  (2) Checklist Fase 0 - levantamento de campo de instrumentacao
Unidades: SNP, MTU, DRD, SDR, BLS, LEM, LRL, SPD
Saida: plano_seguranca/Plano_Seguranca_Matriz_e_Checklist.xlsx
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter

UNIDADES = ["SNP", "MTU", "DRD", "SDR", "BLS", "LEM", "LRL", "SPD"]

# ---- Estilos ----
AZUL = "0B3D5C"
LARANJA = "E86A17"
CINZA_CLARO = "F2F4F6"
VERMELHO = "C0392B"
VERDE = "2E7D32"
AMARELO = "F1C40F"

f_hdr = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
f_tit = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
f_sub = Font(name="Calibri", size=10, italic=True, color="FFFFFF")
f_cell = Font(name="Calibri", size=10, color="333333")
fill_hdr = PatternFill("solid", fgColor=AZUL)
fill_tit = PatternFill("solid", fgColor=AZUL)
fill_alt = PatternFill("solid", fgColor=CINZA_CLARO)
fill_lar = PatternFill("solid", fgColor=LARANJA)
thin = Side(style="thin", color="CCCCCC")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap_top = Alignment(wrap_text=True, vertical="top")
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
ml = Alignment(horizontal="left", vertical="center", wrap_text=True)

wb = Workbook()

# =====================================================================
# ABA 1 — MATRIZ DE RISCO
# =====================================================================
ws = wb.active
ws.title = "Matriz de Risco"

cols1 = [
    ("Unidade", 9),
    ("Área", 14),
    ("Sistema / Equipamento", 20),
    ("Cenário de risco", 26),
    ("Causa raiz", 28),
    ("Probabilidade", 12),
    ("Severidade", 12),
    ("Nível de risco", 12),
    ("Controles existentes", 26),
    ("Ações recomendadas", 32),
    ("Norma", 16),
    ("Prioridade", 11),
    ("Responsável", 14),
    ("Prazo", 12),
    ("Status", 12),
]

# titulo
ncol = len(cols1)
ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
c = ws.cell(1, 1, "MATRIZ DE RISCO — Transportadores por área (grãos, torres, caldeira, destilaria)")
c.font = f_tit
c.fill = fill_tit
c.alignment = ml
ws.row_dimensions[1].height = 24
ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncol)
c = ws.cell(2, 1, "INPASA — Manutenção & Automação  |  Probabilidade/Severidade: Baixa/Média/Alta/Muito Alta  |  "
                  "Risco: CRÍTICO > ALTO > MÉDIO > BAIXO")
c.font = f_sub
c.fill = fill_lar
c.alignment = ml
ws.row_dimensions[2].height = 18

# header
hr = 3
for j, (name, w) in enumerate(cols1, start=1):
    cell = ws.cell(hr, j, name)
    cell.font = f_hdr
    cell.fill = fill_hdr
    cell.alignment = center
    cell.border = border
    ws.column_dimensions[get_column_letter(j)].width = w
ws.row_dimensions[hr].height = 30

# Cenarios padrao por area (preenchidos como ponto de partida)
CENARIOS = [
    # area, sistema, cenario, causa, prob, sev, risco, controles, acoes, norma, prioridade
    ("Grãos", "Elevadores de canecas e correias transportadoras",
     "Incêndio por atrito (rolete/polia travando)",
     "Rolete/mancal travando + correia não-antichama; inspeção visual insuficiente",
     "Alta", "Alta", "CRÍTICO",
     "Termometria de silo; detecção de patinagem parcial",
     "Detecção de patinagem (zero-speed) + temp. de mancal online no PLC com parada segura; correia antichama",
     "NR-12; NFPA 61", "Crítica"),
    ("Grãos", "Elevadores, chutes e pontos de transferência",
     "Explosão/incêndio de poeira combustível",
     "Acúmulo de poeira + fonte de ignição (atrito/fagulha); housekeeping deficiente",
     "Média", "Muito Alta", "CRÍTICO",
     "Limpeza manual periódica",
     "Despoeiramento/housekeeping; detecção de fagulha nos chutes; alívio/contenção de explosão",
     "NFPA 61/654; NR-12", "Crítica"),
    ("Torres", "Recheio (fill) da torre de resfriamento",
     "Ignição do recheio durante manutenção",
     "Respingo de solda cai no recheio em trabalho a quente sem barreira",
     "Média", "Alta", "ALTO",
     "Permissão de trabalho (com falhas)",
     "PT de trabalho a quente com vigia de fogo, isolamento e proteção do recheio; auditoria das PTs",
     "NR-23; NR-12", "Alta"),
    ("Caldeira", "Transportadores de alimentação de biomassa",
     "Incêndio nos transportadores de alimentação",
     "Rolete/correia (mesma natureza dos grãos); material combustível",
     "Média", "Alta", "ALTO",
     "Câmera termográfica anunciando no supervisório",
     "Detecção de patinagem + temp. de mancal + detecção de fagulha; correia antichama",
     "NR-12; NR-13", "Alta"),
    ("Destilaria", "Transportadores/áreas classificadas",
     "Vazamento/ignição em atmosfera explosiva",
     "Falha de equipamento em área classificada (Ex)",
     "Baixa", "Muito Alta", "MÉDIO",
     "Controle mais maduro; equipamentos Ex",
     "Manter conformidade Ex (IEC 60079); inspeção e plano de calibração",
     "IEC 60079", "Média"),
]

r = hr + 1
risco_colors = {
    "CRÍTICO": (VERMELHO, "FFFFFF"),
    "ALTO": (LARANJA, "FFFFFF"),
    "MÉDIO": (AMARELO, "333333"),
    "BAIXO": (VERDE, "FFFFFF"),
}
for u in UNIDADES:
    for cen in CENARIOS:
        area, sis, ce, cau, prob, sev, risco, ctr, aco, norma, prio = cen
        vals = [u, area, sis, ce, cau, prob, sev, risco, ctr, aco, norma, prio, "", "", "Pendente"]
        for j, v in enumerate(vals, start=1):
            cell = ws.cell(r, j, v)
            cell.font = f_cell
            cell.border = border
            cell.alignment = wrap_top if j in (3, 4, 5, 9, 10) else center
            if (r % 2) == 0:
                cell.fill = fill_alt
        # colorir nivel de risco (col 8)
        rc, fc = risco_colors.get(risco, (CINZA_CLARO, "333333"))
        cr = ws.cell(r, 8)
        cr.fill = PatternFill("solid", fgColor=rc)
        cr.font = Font(name="Calibri", size=10, bold=True, color=fc)
        cr.alignment = center
        ws.row_dimensions[r].height = 46
        r += 1

last1 = r - 1
ws.freeze_panes = "A4"
ws.auto_filter.ref = f"A{hr}:{get_column_letter(ncol)}{last1}"

# validacoes (dropdowns)
dv_prob = DataValidation(type="list", formula1='"Baixa,Média,Alta,Muito Alta"', allow_blank=True)
dv_risco = DataValidation(type="list", formula1='"CRÍTICO,ALTO,MÉDIO,BAIXO"', allow_blank=True)
dv_prio = DataValidation(type="list", formula1='"Crítica,Alta,Média,Baixa"', allow_blank=True)
dv_status = DataValidation(type="list", formula1='"Pendente,Em andamento,Concluído"', allow_blank=True)
ws.add_data_validation(dv_prob); ws.add_data_validation(dv_risco)
ws.add_data_validation(dv_prio); ws.add_data_validation(dv_status)
dv_prob.add(f"F4:G{last1}")
dv_risco.add(f"H4:H{last1}")
dv_prio.add(f"L4:L{last1}")
dv_status.add(f"O4:O{last1}")

# =====================================================================
# ABA 2 — CHECKLIST FASE 0
# =====================================================================
ws2 = wb.create_sheet("Checklist Fase 0")
cols2 = [
    ("ID", 6),
    ("Unidade", 9),
    ("Área", 12),
    ("TAG do equipamento", 16),
    ("Tipo", 14),
    ("Detecção de patinagem", 12),
    ("Sensor temp. de mancal", 12),
    ("Detecção de fagulha", 12),
    ("Sensor de alinhamento", 12),
    ("Correia antichama", 12),
    ("Parada de emerg. (cabo)", 12),
    ("Integrado ao PLC/SCADA", 12),
    ("Estado de conservação", 14),
    ("Ação requerida", 28),
    ("Prioridade", 11),
    ("Inspetor", 14),
    ("Data", 12),
]
ncol2 = len(cols2)
ws2.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol2)
c = ws2.cell(1, 1, "CHECKLIST DE LEVANTAMENTO DE CAMPO — FASE 0 (inventário de instrumentação dos transportadores)")
c.font = f_tit; c.fill = fill_tit; c.alignment = ml
ws2.row_dimensions[1].height = 24
ws2.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncol2)
c = ws2.cell(2, 1, "Preencher em campo por equipamento (elevador/correia). Colunas Sim/Não/N/A com lista suspensa. "
                   "Objetivo: mapear gaps de patinagem, temperatura de mancal e detecção de fagulha.")
c.font = f_sub; c.fill = fill_lar; c.alignment = ml
ws2.row_dimensions[2].height = 18

hr2 = 3
for j, (name, w) in enumerate(cols2, start=1):
    cell = ws2.cell(hr2, j, name)
    cell.font = f_hdr; cell.fill = fill_hdr; cell.alignment = center; cell.border = border
    ws2.column_dimensions[get_column_letter(j)].width = w
ws2.row_dimensions[hr2].height = 34

# linhas em branco para preenchimento: 12 por unidade
LINHAS_POR_UNIDADE = 12
r = hr2 + 1
idx = 1
for u in UNIDADES:
    for _ in range(LINHAS_POR_UNIDADE):
        vals = [idx, u] + [""] * (ncol2 - 2)
        for j, v in enumerate(vals, start=1):
            cell = ws2.cell(r, j, v)
            cell.font = f_cell; cell.border = border
            cell.alignment = wrap_top if j == 14 else center
            if (r % 2) == 0:
                cell.fill = fill_alt
        ws2.row_dimensions[r].height = 22
        r += 1
        idx += 1

last2 = r - 1
ws2.freeze_panes = "A4"
ws2.auto_filter.ref = f"A{hr2}:{get_column_letter(ncol2)}{last2}"

# validacoes
dv_sn = DataValidation(type="list", formula1='"Sim,Não,N/A"', allow_blank=True)
dv_tipo = DataValidation(type="list", formula1='"Elevador de canecas,Correia transportadora,Redler,Rosca,Outro"', allow_blank=True)
dv_estado = DataValidation(type="list", formula1='"Bom,Regular,Ruim"', allow_blank=True)
dv_prio2 = DataValidation(type="list", formula1='"Crítica,Alta,Média,Baixa"', allow_blank=True)
for dv in (dv_sn, dv_tipo, dv_estado, dv_prio2):
    ws2.add_data_validation(dv)
dv_tipo.add(f"E4:E{last2}")
for col in ["F", "G", "H", "I", "J", "K", "L"]:
    dv_sn.add(f"{col}4:{col}{last2}")
dv_estado.add(f"M4:M{last2}")
dv_prio2.add(f"O4:O{last2}")

# realce condicional: "Não" em vermelho nas colunas de instrumentacao (gaps)
red_fill = PatternFill("solid", fgColor="F8D7DA")
red_font = Font(color="C0392B", bold=True)
for col in ["F", "G", "H", "I", "J", "K", "L"]:
    ws2.conditional_formatting.add(
        f"{col}4:{col}{last2}",
        CellIsRule(operator="equal", formula=['"Não"'], fill=red_fill, font=red_font))

out = "/home/user/ICG2025/plano_seguranca/Plano_Seguranca_Matriz_e_Checklist.xlsx"
wb.save(out)
print("OK:", out)
print("Aba1 linhas de risco:", last1 - hr, "| Aba2 linhas checklist:", last2 - hr2)
