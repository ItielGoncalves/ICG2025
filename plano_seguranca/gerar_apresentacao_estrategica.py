# -*- coding: utf-8 -*-
"""
Apresentacao ESTRATEGICA (CEO / VPs) — Plano de Substituicao de Correias.
Converte o dashboard operacional em narrativa de diretoria.
Slide 2 (apos capa): quando cada unidade estara regularizada.
Saida: plano_seguranca/Apresentacao_Estrategica_Correias_INPASA.pptx
"""
from datetime import date
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ---------- DADOS ----------
raw = [
 ('NVM','MILHO','TRL-519004-4','09/05/2026'),('NVM','MILHO','TRL-519004-1','20/08/2026'),
 ('NVM','MILHO','TRL-519004-2','25/08/2026'),('NVM','MILHO','TRL-519004-3','30/08/2026'),
 ('NVM','MILHO','TRL-519004-5','10/09/2026'),('NVM','MILHO','TRL-519424','05/08/2026'),
 ('NVM','MILHO','TRL-519426','30/09/2026'),('NVM','MILHO','TRL-519425','30/09/2026'),
 ('NVM','MILHO','TRL-719102','30/09/2026'),('NVM','MILHO','TRL-719104','20/09/2026'),
 ('NVM','MILHO','TRL-719103','05/10/2026'),('NVM','MILHO','TRL-719101','30/09/2026'),
 ('NVM','DDGS','TRL-509041','10/10/2026'),('NVM','MILHO','TRL-519004-7','01/08/2026'),
 ('NVM','MILHO','TRL-519004-6','30/09/2026'),
 ('SNP','BIOMASSA','TRL-401002A','01/08/2026'),('SNP','BIOMASSA','TRL-401002B','01/08/2026'),
 ('SNP','BIOMASSA','TRL-631501','17/07/2026'),('SNP','BIOMASSA','TRL-631504','17/07/2026'),
 ('SNP','BIOMASSA','TRL-631512','17/07/2026'),('SNP','DDGS','TRL-409302','30/09/2026'),
 ('SNP','GRAOS','TRL-419419SA-1','30/09/2026'),('SNP','GRAOS','TRL-419419SA-2','30/09/2026'),
 ('SNP','GRAOS','TRL-419419SB-2','30/09/2026'),('SNP','GRAOS','TRL-419SB3','30/09/2026'),
 ('SNP','GRAOS','TRL-419419SB-5','30/09/2026'),('SNP','GRAOS','TRL-419419SB-6','30/09/2026'),
 ('SNP','GRAOS','TRL-619038','30/09/2026'),('SNP','GRAOS','TRL-619023','30/09/2026'),
 ('SNP','GRAOS','TRL-619041','30/09/2026'),('SNP','GRAOS','TRL-1011235','30/09/2026'),
 ('SDL','DDGS','TRL-1209011','17/07/2026'),('SDL','GRAOS','TRL-1219523','30/09/2026'),
 ('BLS','GRAOS','TRL-1519203A','SUBSTITUÍDA'),('BLS','DDGS','TRL-1509208','17/07/2026'),
 ('BLS','DDGS','TRL-1509209','17/07/2026'),('BLS','DDGS','TRL-1509403','17/07/2026'),
 ('BLS','ALMOX','ESTOQUE','30/09/2026'),
 ('SNP','BIOMASSA','TRL-1031004','30/09/2026'),('SNP','BIOMASSA','TRL-1031008','SUBSTITUÍDA'),
 ('SNP','BIOMASSA','TRL-1031009','30/09/2026'),('SNP','BIOMASSA','TRL-1031011','30/09/2026'),
 ('SNP','BIOMASSA','TRL-1031012','30/09/2026'),
]
TODAY = date(2026, 7, 2)
MESES = {1:'jan',2:'fev',3:'mar',4:'abr',5:'mai',6:'jun',7:'jul',8:'ago',9:'set',10:'out',11:'nov',12:'dez'}

def parse(d):
    if d == 'SUBSTITUÍDA':
        return None
    dd, mm, yy = d.split('/')
    return date(int(yy), int(mm), int(dd))

U = {}
for u, a, t, d in raw:
    U.setdefault(u, []).append((t, a, parse(d), d))

def stats(u):
    rows = U[u]
    total = len(rows)
    subst = sum(1 for _, _, dt, _ in rows if dt is None)
    estoque = sum(1 for t, _, _, _ in rows if t == 'ESTOQUE')
    venc = sum(1 for _, _, dt, _ in rows if dt and dt < TODAY)
    urg = sum(1 for _, _, dt, _ in rows if dt and 0 <= (dt - TODAY).days <= 30)
    pend = [dt for _, _, dt, _ in rows if dt]
    return dict(total=total, subst=subst, estoque=estoque, venc=venc, urg=urg,
                pend=total - subst, first=min(pend), last=max(pend))

NOMES = {'NVM': 'Nova Mutum (NVM)', 'SNP': 'Sinop (SNP)', 'SDL': 'Sidrolândia (SDL)', 'BLS': 'Balsas (BLS)'}
CORES = {'NVM': RGBColor(0x3B,0x82,0xF6), 'SNP': RGBColor(0xF9,0x73,0x16),
         'SDL': RGBColor(0x22,0xC5,0x5E), 'BLS': RGBColor(0xA8,0x55,0xF7)}
ORDEM = ['NVM', 'SNP', 'SDL', 'BLS']
ST = {u: stats(u) for u in ORDEM}
GERAL_LAST = max(s['last'] for s in ST.values())
TOTAL = sum(s['total'] for s in ST.values())
SUBST = sum(s['subst'] for s in ST.values())
URG = sum(s['urg'] for s in ST.values())
VENC = sum(s['venc'] for s in ST.values())

def fdata(d):
    return d.strftime('%d/%m/%Y')

# ---------- ESTILO ----------
BG = RGBColor(0x0F,0x0F,0x1A)
CARD = RGBColor(0x1E,0x1E,0x30)
CARD2 = RGBColor(0x25,0x25,0x38)
LARANJA = RGBColor(0xF9,0x73,0x16)
BRANCO = RGBColor(0xFF,0xFF,0xFF)
TXT = RGBColor(0xE0,0xE0,0xE0)
MUT = RGBColor(0x9C,0xA3,0xAF)
VERDE = RGBColor(0x22,0xC5,0x5E)
VERM = RGBColor(0xEF,0x44,0x44)
AMAR = RGBColor(0xEA,0xB3,0x08)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]

def slide():
    s = prs.slides.add_slide(BLANK)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    bg.fill.solid(); bg.fill.fore_color.rgb = BG
    bg.line.fill.background(); bg.shadow.inherit = False
    return s

def rect(s, x, y, w, h, color, line=None, line_w=1.0, rounded=False):
    shape = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    sp = s.shapes.add_shape(shape, x, y, w, h)
    sp.fill.solid(); sp.fill.fore_color.rgb = color
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line; sp.line.width = Pt(line_w)
    sp.shadow.inherit = False
    return sp

def txt(s, x, y, w, h, text, size=18, bold=False, color=TXT,
        align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font='Inter', spacing=None):
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = Pt(2); tf.margin_right = Pt(2)
    for i, ln in enumerate(text.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if spacing:
            p.line_spacing = spacing
        r = p.add_run(); r.text = ln
        r.font.size = Pt(size); r.font.bold = bold
        r.font.color.rgb = color; r.font.name = font
    return tb

def head(s, kicker, title):
    rect(s, 0, 0, SW, Inches(1.2), CARD)
    rect(s, 0, Inches(1.2), SW, Inches(0.04), LARANJA)
    txt(s, Inches(0.55), Inches(0.16), Inches(12), Inches(0.3),
        kicker, size=12, bold=True, color=LARANJA)
    txt(s, Inches(0.55), Inches(0.46), Inches(12.2), Inches(0.6),
        title, size=25, bold=True, color=BRANCO, anchor=MSO_ANCHOR.MIDDLE)

def foot(s, n):
    txt(s, Inches(0.55), Inches(7.05), Inches(9), Inches(0.35),
        "INPASA  •  Plano de Substituição de Correias Transportadoras  •  2026",
        size=9, color=RGBColor(0x6B,0x72,0x80))
    txt(s, Inches(12.2), Inches(7.05), Inches(0.9), Inches(0.35),
        str(n), size=9, color=RGBColor(0x6B,0x72,0x80), align=PP_ALIGN.RIGHT)

# =====================================================================
# 1 — CAPA
# =====================================================================
s = slide()
rect(s, 0, Inches(4.55), SW, Inches(0.05), LARANJA)
txt(s, Inches(0.9), Inches(1.5), Inches(11.5), Inches(0.5),
    "🏭  INPASA  •  APRESENTAÇÃO À DIRETORIA", size=15, bold=True, color=LARANJA)
txt(s, Inches(0.9), Inches(2.25), Inches(11.6), Inches(1.8),
    "Plano de Substituição de\nCorreias Transportadoras", size=44, bold=True, color=BRANCO,
    spacing=1.05)
txt(s, Inches(0.9), Inches(4.75), Inches(11.6), Inches(1.0),
    "Status do programa e cronograma de regularização por unidade\n"
    "Nova Mutum • Sinop • Sidrolândia • Balsas", size=17, color=MUT, spacing=1.2)
txt(s, Inches(0.9), Inches(6.55), Inches(11.6), Inches(0.5),
    "Gestão de Manutenção Industrial  |  Data de referência: 02/07/2026", size=12,
    color=RGBColor(0x6B,0x72,0x80))

# =====================================================================
# 2 — REGULARIZACAO POR UNIDADE  (PAGINA PEDIDA)
# =====================================================================
s = slide()
head(s, "A PERGUNTA DA DIRETORIA", "Quando cada unidade estará regularizada?")
# headline
rect(s, Inches(0.55), Inches(1.45), Inches(12.25), Inches(0.85), CARD2,
     line=LARANJA, line_w=1.25, rounded=True)
txt(s, Inches(0.8), Inches(1.45), Inches(11.8), Inches(0.85),
    f"Programa 100% regularizado até  {fdata(GERAL_LAST)}   "
    f"|   {TOTAL} correias no plano  •  {SUBST} já substituídas  •  {URG} urgentes (≤30 dias)"
    + (f"  •  ⚠ {VENC} vencida" if VENC else ""),
    size=15, bold=True, color=BRANCO, anchor=MSO_ANCHOR.MIDDLE)

# 4 cards
cx = Inches(0.55); cw = Inches(2.95); gap = Inches(0.13); cy = Inches(2.55); ch = Inches(3.15)
frases = {
 'NVM': "São {t} esteiras no plano. A última será concluída em {last} — data de regularização da unidade.",
 'SNP': "São {t} esteiras no plano; {sub} já substituída. A última será concluída em {last}.",
 'SDL': "São {t} esteiras no plano. Ambas concluídas até {last} — unidade regularizada nessa data.",
 'BLS': "São 4 esteiras + 1 sobressalente em estoque; {sub} já substituída. Última em {last}.",
}
for i, u in enumerate(ORDEM):
    st = ST[u]; x = Emu(int(cx) + i * (int(cw) + int(gap)))
    rect(s, x, cy, cw, ch, CARD, rounded=True)
    rect(s, x, cy, cw, Inches(0.12), CORES[u])  # top accent
    txt(s, Emu(int(x)+Inches(0.2)), Emu(int(cy)+Inches(0.28)), Emu(int(cw)-Inches(0.4)), Inches(0.4),
        u, size=13, bold=True, color=CORES[u])
    txt(s, Emu(int(x)+Inches(0.2)), Emu(int(cy)+Inches(0.62)), Emu(int(cw)-Inches(0.4)), Inches(0.5),
        NOMES[u].split('(')[0].strip(), size=15, bold=True, color=BRANCO)
    txt(s, Emu(int(x)+Inches(0.2)), Emu(int(cy)+Inches(1.15)), Emu(int(cw)-Inches(0.4)), Inches(0.35),
        "REGULARIZADA EM", size=10, bold=True, color=MUT)
    txt(s, Emu(int(x)+Inches(0.2)), Emu(int(cy)+Inches(1.42)), Emu(int(cw)-Inches(0.4)), Inches(0.55),
        fdata(st['last']), size=22, bold=True, color=CORES[u])
    frase = frases[u].format(t=st['total'], sub=st['subst'], last=fdata(st['last']))
    txt(s, Emu(int(x)+Inches(0.2)), Emu(int(cy)+Inches(2.05)), Emu(int(cw)-Inches(0.4)), Inches(1.0),
        frase, size=10.5, color=TXT, spacing=1.05)
    if st['venc']:
        txt(s, Emu(int(x)+Inches(0.2)), Emu(int(cy)+Inches(2.78)), Emu(int(cw)-Inches(0.4)), Inches(0.3),
            f"⚠ {st['venc']} vencida (atrasada)", size=10, bold=True, color=VERM)
foot(s, 2)

# =====================================================================
# 3 — CRONOGRAMA (marcos de regularizacao)
# =====================================================================
s = slide()
head(s, "LINHA DO TEMPO", "Marcos de regularização (jul → out/2026)")
tl_x = Inches(2.6); tl_w = Inches(9.9)
start = date(2026,7,1); end = date(2026,10,31)
span = (end - start).days
months = [('Julho',7),('Agosto',8),('Setembro',9),('Outubro',10)]
top = Inches(1.75)
# grade de meses
for name, m in months:
    md = date(2026,m,1)
    px = int(tl_x) + int(((md-start).days/span) * int(tl_w))
    rect(s, Emu(px), Emu(int(top)), Pt(1), Inches(4.4), RGBColor(0x2A,0x2A,0x40))
    txt(s, Emu(px), Emu(int(top)-Inches(0.02)), Inches(1.8), Inches(0.3),
        name, size=12, bold=True, color=LARANJA)
# linha HOJE
hx = int(tl_x) + int(((TODAY-start).days/span) * int(tl_w))
rect(s, Emu(hx), Emu(int(top)+Inches(0.3)), Pt(2), Inches(4.05), VERM)
txt(s, Emu(hx-Inches(0.35)), Emu(int(top)+Inches(0.02)), Inches(1.2), Inches(0.28),
    "HOJE", size=10, bold=True, color=VERM)
# barras por unidade (do inicio ate a data de regularizacao)
rowh = Inches(0.95); ry = Emu(int(top)+Inches(0.55))
for u in ORDEM:
    st = ST[u]
    txt(s, Inches(0.55), ry, Inches(1.95), rowh, f"{u}\n{st['total']} correias",
        size=12, bold=True, color=CORES[u], anchor=MSO_ANCHOR.MIDDLE)
    x0 = int(tl_x)
    x1 = int(tl_x) + int(((st['last']-start).days/span) * int(tl_w))
    bar = rect(s, Emu(x0), Emu(int(ry)+Inches(0.22)), Emu(max(x1-x0, int(Inches(0.2)))),
               Inches(0.42), CORES[u], rounded=True)
    txt(s, Emu(x1+Inches(0.1)), Emu(int(ry)+Inches(0.2)), Inches(1.8), Inches(0.45),
        fdata(st['last']), size=11, bold=True, color=BRANCO, anchor=MSO_ANCHOR.MIDDLE)
    ry = Emu(int(ry) + int(rowh))
txt(s, Inches(0.55), Inches(6.5), Inches(12), Inches(0.4),
    f"Conclusão do programa: {fdata(GERAL_LAST)} (última correia — Nova Mutum / DDGS).",
    size=13, bold=True, color=LARANJA)
foot(s, 3)

# =====================================================================
# 4 — DETALHE POR UNIDADE (tabela resumo)
# =====================================================================
s = slide()
head(s, "VISÃO CONSOLIDADA", "Resumo executivo por unidade")
cols = ["Unidade", "Correias\nno plano", "Já\nsubstituídas", "Em\nexecução",
        "Próxima\nsubstituição", "Regularização\n(última)", "Alerta"]
colw = [0.20, 0.12, 0.12, 0.11, 0.15, 0.15, 0.15]
tx = Inches(0.55); tw = Inches(12.25); hy = Inches(1.55); rh = Inches(0.55)
# header
xx = int(tx)
for j, c in enumerate(cols):
    w = int(tw * colw[j])
    rect(s, Emu(xx), Emu(int(hy)), Emu(w), Inches(0.7), CARD2)
    txt(s, Emu(xx), Emu(int(hy)), Emu(w), Inches(0.7), c, size=11, bold=True,
        color=LARANJA, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    xx += w
ry = int(hy) + int(Inches(0.7))
for u in ORDEM:
    st = ST[u]
    alerta = (f"{st['venc']} vencida" if st['venc'] else
              (f"{st['urg']} urgentes" if st['urg'] else "—"))
    acor = VERM if st['venc'] else (AMAR if st['urg'] else MUT)
    vals = [NOMES[u], str(st['total']), str(st['subst']), str(st['pend']),
            fdata(st['first']), fdata(st['last']), alerta]
    xx = int(tx)
    for j, v in enumerate(vals):
        w = int(tw * colw[j])
        rect(s, Emu(xx), Emu(ry), Emu(w), rh, CARD if ORDEM.index(u) % 2 == 0 else BG)
        col = CORES[u] if j == 0 else (acor if j == 6 else TXT)
        b = (j in (0, 5, 6))
        txt(s, Emu(xx), Emu(ry), Emu(w), rh, v, size=11, bold=b, color=col,
            align=PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        xx += w
    ry += int(rh)
# linha total
xx = int(tx)
tot = [f"TOTAL ({len(ORDEM)} unidades)", str(TOTAL), str(SUBST), str(TOTAL-SUBST),
       "—", fdata(GERAL_LAST), (f"{VENC} venc / {URG} urg")]
for j, v in enumerate(tot):
    w = int(tw * colw[j])
    rect(s, Emu(xx), Emu(ry), Emu(w), rh, LARANJA)
    txt(s, Emu(xx), Emu(ry), Emu(w), rh, v, size=11, bold=True, color=BG,
        align=PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    xx += w
foot(s, 4)

# =====================================================================
# 5 — PONTOS DE ATENCAO E PROXIMOS PASSOS
# =====================================================================
s = slide()
head(s, "DECISÃO", "Pontos de atenção e próximos passos")
# coluna riscos
rect(s, Inches(0.55), Inches(1.55), Inches(6.0), Inches(4.9), CARD, rounded=True)
txt(s, Inches(0.8), Inches(1.75), Inches(5.5), Inches(0.4), "⚠  Riscos / prioridades",
    size=16, bold=True, color=VERM)
riscos = [
    f"NVM: 1 correia VENCIDA (atrasada desde 09/05/2026) — regularizar imediatamente.",
    f"{URG} correias urgentes (≤30 dias), concentradas em 17/07 e 01/08 — biomassa (SNP), DDGS (SDL/BLS).",
    "Áreas de biomassa e grãos: maior exposição a incêndio por atrito (rolete/correia).",
    "Conclusão do programa depende de NVM (última em 10/10/2026).",
]
ty = Inches(2.3)
for r in riscos:
    txt(s, Inches(0.85), ty, Inches(5.5), Inches(1.0), "•  " + r, size=13, color=TXT, spacing=1.05)
    ty = Emu(int(ty) + int(Inches(1.02)))
# coluna proximos passos
rect(s, Inches(6.8), Inches(1.55), Inches(6.0), Inches(4.9), CARD, rounded=True)
txt(s, Inches(7.05), Inches(1.75), Inches(5.5), Inches(0.4), "✓  Próximos passos",
    size=16, bold=True, color=VERDE)
passos = [
    "Priorizar a correia vencida de NVM e as urgentes de 17/07 nesta semana.",
    "Confirmar disponibilidade de correia antichama e mão de obra para o pico de set/2026.",
    "Reportar avanço quinzenal por unidade (esteiras concluídas x plano).",
    "Antecipar itens de NVM para fechar o programa antes de 10/10/2026, se possível.",
]
ty = Inches(2.3)
for p in passos:
    txt(s, Inches(7.1), ty, Inches(5.5), Inches(1.0), "•  " + p, size=13, color=TXT, spacing=1.05)
    ty = Emu(int(ty) + int(Inches(1.02)))
foot(s, 5)

out = "/home/user/ICG2025/plano_seguranca/Apresentacao_Estrategica_Correias_INPASA.pptx"
prs.save(out)
print("OK:", out, "| slides:", len(prs.slides._sldIdLst))
